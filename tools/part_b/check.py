"""Verify Part B content, source-linked evidence and fresh rendered pages."""
from pathlib import Path
import hashlib,json,re,zipfile,xml.etree.ElementTree as ET
from docx import Document
from docx.oxml.ns import qn
from PIL import ImageFont
from content import CASES,INTRO,MODULES
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'docs/part-b';QA=ROOT/'backend/target/part-b-qa'
checks=[]
def check(name,ok):
    checks.append({'name':name,'passed':bool(ok)})
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def run():
    path=OUT/'PartB_Submission.docx';d=Document(path);template=ROOT/'docs/templates/PartB-Submission-Template.docx'
    check('Retained template bytes',sha(template)=='4e3f451e8ed369ab857418437e2cd1f41b27cf52273bb9f26ce4aece4cb5afc8')
    original=Document(template)
    check('Section count',len(d.sections)==len(original.sections))
    for i,(a,b) in enumerate(zip(d.sections,original.sections)):
        check('Page geometry '+str(i),all(getattr(a,k)==getattr(b,k) for k in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin','header_distance','footer_distance']))
    body='\n'.join(''.join(node.xpath('.//w:t/text()')) for node in d.element.body.xpath('.//w:p'))
    headings=[p.text for p in d.paragraphs if p.style.name.startswith('Heading')]
    expected=['Part B. 功能原型','1. PDF Extractor 项目','1.1 项目介绍','1.2 交付物 2 应用设计','1.2.1 系统架构','1.2.2 类与模块设计','1.2.3 界面设计 WND 与 Swagger 文档','1.3 交付物 3 应用软件','1.3.1 可执行软件文件','1.4 交付物 4 测试','1.4.1 JUnit 测试及执行记录','1.4.2 可测量验收标准','1.5 交付物 5 源码质量','1.6 交付物 6 范围边界','1.7 交付物 7 课堂演示']
    check('Template section order',headings==expected)
    check('Chinese metadata',d.core_properties.language=='zh-CN')
    check('Three paragraph introduction',len(INTRO)==3 and all(x in body for x in INTRO))
    check('Chinese member headings','成员英文姓名' in d.tables[0].cell(0,1).text and '学号' in d.tables[0].cell(0,2).text)
    check('Five unfilled member rows',len(d.tables[0].rows)==6 and all(not c.text.strip() for row in d.tables[0].rows[1:] for c in row.cells[1:]))
    for phrase in ['here you should','here include','AI test agents','按你的要求','待用户确认','不能冒充','不代表已经实现','登录模块','助手模块','Error! Reference','错误!']:
        check('No instructional residue '+phrase,phrase not in body)
    check('Native TOC retained',any('TOC' in s for s in d.element.body.xpath('.//w:instrText/text()')))
    check('Five figures',len(d.inline_shapes)==5)
    check('Actual synchronous processing stated','PDF 解析及输出生成仍位于 Web JVM 内' in body and '基础 11 项 PDF 操作不入消息队列' in body)
    check('Middleware roles stated',all(s in body for s in ['Nginx','PostgreSQL','Redis','RabbitMQ','Elasticsearch','S3']))
    check('No completed isolated PDF process claim','受限 Java PDF 子进程' not in body)
    check('Source quality template columns',any(all(x in ' '.join(c.text for c in t.rows[0].cells) for x in ['Requirement','Implementation','Test Case','Evidence']) for t in d.tables))
    report=ET.parse(OUT/'evidence/core-junit.xml').getroot();cases=report.findall('.//testcase')
    check('Actual JUnit tests',len(cases)==93 and len(report.findall('testsuite'))==11)
    check('No environment properties in public report',not report.findall('.//properties') and not report.findall('.//system-out'))
    check('No failed or skipped cases',not any(report.findall('.//'+tag) for tag in ['failure','error','skipped']))
    manifest=json.loads((OUT/'evidence/manifest.json').read_text('utf-8'))
    for item in manifest['files']:
        p=OUT/'evidence'/item['path'];check('Evidence checksum '+item['path'],p.is_file() and sha(p)==item['sha256'])
    for item in manifest['source_files']:
        p=ROOT/item['path'];check('Source checksum '+item['path'],p.is_file() and sha(p)==item['sha256'])
    for item in manifest['suite_sources']:
        p=ROOT/('backend/target/surefire-reports/TEST-cn.edu.pdfworkbench.'+item['suite']+'.xml')
        check('Execution report provenance '+item['suite'],p.is_file() and sha(p)==item['source_report_sha256'])
    check('Browser summary',re.search(r'8 passed', (OUT/'evidence/browser-summary.txt').read_text('utf-8')))
    for id,title,cls,method,*rest in CASES:
        matches=[c for c in cases if c.get('classname')=='cn.edu.pdfworkbench.'+cls and c.get('name').startswith(method+'(')]
        # Surefire names include parameter signatures for parameterized instances.
        if not matches:matches=[c for c in cases if c.get('classname')=='cn.edu.pdfworkbench.'+cls and c.get('name')==method]
        check('Executed method '+id,bool(matches))
        if id=='TC05':check('Sixteen permission combinations',len(matches)==16)
        check('Test record '+id,id+' '+title in body and method in body and all(s in body for s in rest[:3]))
        check('Test source method '+id,method+'(' in (ROOT/f'backend/src/test/java/cn/edu/pdfworkbench/{cls}.java').read_text('utf-8'))
        for file in rest[-1].split('；'):check('Case evidence file '+id+' '+file,(OUT/'evidence'/file).is_file())
    check('At least five complete records',len(CASES)>=5 and len([t for t in d.tables if t.cell(0,0).text=='记录项'])==len(CASES))
    for group in ['FR1','FR2','FR3','FR4','FR5','FR6','FR7','FR9','FR10','FR11','FR12','FR13–FR16','FR17–FR19','WM01','WM02']:
        check('Requirement coverage '+group,group in body)
    check('Optional attachment boundary','FR8 附件提取不在范围' in body)
    check('Fourteen records / evidence counts',manifest['junit_tests']==93 and manifest['browser_tests']==8 and '14 条完整记录' in body)
    for name in ['architecture','modules','wnd']:
        fig=json.loads((OUT/'diagrams'/f'{name}.json').read_text('utf-8'));ids={x['id'] for x in fig['nodes']}
        check('Diagram endpoints '+name,all(e['from'] in ids and e['to'] in ids for e in fig['edges']))
        font=ImageFont.truetype('C:/Windows/Fonts/simhei.ttf',30)
        check('Node text fits '+name,all(max(font.getlength(line) for line in n['label'].split('\n'))<n['width']-12 for n in fig['nodes']))
    wnd=json.loads((OUT/'diagrams/wnd.json').read_text('utf-8'))
    check('WND parity with Part A',set(n['id'] for n in wnd['nodes'])=={'W'+str(i).zfill(2) for i in range(11)})
    check('Retry navigation',any(e['from']=='W10' and e['to']=='W01' for e in wnd['edges']))
    render=json.loads((QA/'render.json').read_text('utf-8-sig'));raster=json.loads((QA/'raster.json').read_text('utf-8'));pages=json.loads((QA/'page-index.json').read_text('utf-8'))
    check('Current Word rendered',render['WordSHA256'].lower()==sha(path))
    check('Current PDF rasterized',render['PdfSHA256'].lower()==raster['pdf_sha256']==sha(QA/'PartB.pdf'))
    check('Page count agreement',render['Pages']==raster['pages']==len(pages))
    for page in pages:
        text=(QA/'pages'/f'page-{page["page"]}.txt').read_text('utf-8')
        content=re.sub(r'Introduction to Software Engineering Course.*?Page \d+','',text,flags=re.S).strip()
        screenshot_page='图 4 文本参数' in content and '图 5 表格' in content
        check('Nonblank page '+str(page['page']),len(content)>60 or screenshot_page)
        check('Page image '+str(page['page']),(QA/'pages'/f'page-{page["page"]}.png').is_file())
    with zipfile.ZipFile(template) as a,zipfile.ZipFile(path) as b:
        for part in [n for n in a.namelist() if n.startswith('word/theme/')]:check('Preserved theme '+part,a.read(part)==b.read(part))
    result={'checks':len(checks),'passed':sum(c['passed'] for c in checks),'failures':[c['name'] for c in checks if not c['passed']],'pages':len(pages),'word_sha256':sha(path)}
    (QA/'checks.json').write_text(json.dumps({'summary':result,'checks':checks},ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False));return not result['failures']
if __name__=='__main__':raise SystemExit(0 if run() else 1)
