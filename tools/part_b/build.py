"""Build the Chinese functional prototype report from the retained template."""
from pathlib import Path
import sys,json,hashlib
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).resolve().parents[2]
sys.path.append(str(ROOT/'tools/part_a'))
import submission as base
import graphics
from content import INTRO,MODULES,CASES,ACCEPTANCE,SCOPE,DEMO
OUT=ROOT/'docs/part-b';QA=ROOT/'backend/target/part-b-qa'
TEMPLATE=ROOT/'docs/templates/PartB-Submission-Template.docx'
TEMPLATE_SHA='4e3f451e8ed369ab857418437e2cd1f41b27cf52273bb9f26ce4aece4cb5afc8'
def p(d,text):return d.add_paragraph(text)
def h(d,text,level=2,page=False):return base.h(d,text,level,page)
def table(d,headers,rows,widths):
    base.table(d,headers,rows,widths)
    spacer=d.paragraphs[-1]
    spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.space_after=Pt(3);spacer.paragraph_format.line_spacing=1
    spacer.add_run(' ').font.size=Pt(1)
    for row in d.tables[-1].rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size=Pt(10)
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
def figure(d,path,caption):
    para=p(d,'');para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.keep_with_next=True
    para.add_run().add_picture(str(path),width=Cm(15.9))
    cap=p(d,caption);cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:run.font.size=Pt(10.5)
def diagrams():
    out=OUT/'diagrams';out.mkdir(parents=True,exist_ok=True);items=[]
    f=graphics.Figure('architecture','基础 PDF 请求与本机部署',1400,1040)
    f.node('browser','浏览器工作台\nVue 3 / TypeScript / PDF.js',700,95,1080,130,'external')
    f.node('nginx','Nginx 同源入口\n静态构建文件 / API 反向代理',700,310,1080,130)
    f.node('web','Windows 原生 Java Web\n控制器 → PdfService → PDF 工具库',700,555,1080,160)
    f.node('temp','每请求临时目录\n输入副本及处理结果',360,815,550,140,'store')
    f.node('return','JSON / CSV / ZIP / PDF\n浏览器查看与另存',1040,815,550,140,'external')
    f.edge('browser','nginx','选择文件及提交参数',[(700,160),(700,245)],(700,200),27)
    f.edge('nginx','web','/api/v1/pdf/*',[(700,375),(700,475)],(700,425),27)
    f.edge('web','temp','校验后处理',[(360,635),(360,745)],(360,690),27)
    f.edge('web','return','响应交付',[(1040,635),(1040,745)],(1040,690),27)
    f.text(700,970,'基础 PDF 运算在 Web 进程内同步执行；结束后清理请求目录。',28)
    items.append(f.save(out))
    f=graphics.Figure('modules','控制器 服务与工具模块',1400,1050)
    f.node('controller','ExtractionController / ModificationController\nSignatureController',700,90,1220,130)
    f.node('service','PdfService\n配额 · 输入检查 · 临时目录 · 输出交付',700,330,920,140)
    f.node('extract','PdfTextUtils / PdfInfoUtils\nPdfTableUtils / PdfImageUtils\n文本 网格和图片收集器',360,610,620,170)
    f.node('modify','PdfEncryptionUtils\nPdfWatermarkUtils / PdfSignatureUtils\n签名验证与材料加载',1040,610,620,170)
    f.node('shared','PageSelection / Rect / 选项与结果模型 / PdfFiles',700,875,1250,100)
    f.edge('controller','service','选项对象与请求文件',[(700,155),(700,260)],(700,205),27)
    for ident,x in [('extract',360),('modify',1040)]:
        f.edge('service',ident,'协调处理',[(700,400),(700,460),(x,460),(x,525)],(x,470),26)
        f.edge(ident,'shared','模型及文件操作',[(x,695),(x,825)],(x,760),26)
    f.text(700,995,'异常统一交由 ApiExceptionHandler 转换为 HTTP 状态与错误体。',27)
    items.append(f.save(out))
    f=graphics.Figure('wnd','工作台页面导航',1400,1180)
    f.node('W00','W00 选择 PDF 文件',700,75,700,100,'external')
    f.node('W01','W01 文档预览及页码导航',700,260,700,110)
    panels=['W02 文本与检索','W03 表格与 CSV','W04 图片与 ZIP','W05 文档信息','W06 加密及权限','W07 数字签名','W08 文字与图片水印']
    for i,label in enumerate(panels):
        y=420+i*90;ident=label.split()[0];f.node(ident,label,410,y,520,65)
        f.edge('W01',ident,'',[(350,260),(70,260),(70,y),(150,y)])
        f.edge(ident,'W09','',[(670,y),(780,y),(780,600),(830,600)])
    f.node('W09','W09 检查结果\n保存支持的输出',1080,600,500,140)
    f.node('W10','W10 状态及错误反馈\n修正参数后重试',1080,890,500,140)
    f.edge('W00','W01','加载有效文档',[(700,125),(700,205)],(700,165),25)
    f.edge('W09','W10','操作失败',[(1080,670),(1080,820)],(1080,745),25)
    f.edge('W10','W01','修正后返回当前文档',[(1080,960),(1080,1070),(25,1070),(25,260),(350,260)],(850,1070),25)
    items.append(f.save(out))
    (OUT/'diagrams/manifest.json').write_text(json.dumps(items,ensure_ascii=False,indent=2),encoding='utf-8')
def run():
    assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()==TEMPLATE_SHA
    evidence=json.loads((OUT/'evidence/manifest.json').read_text('utf-8'))
    assert evidence['junit_tests']==93 and evidence['browser_tests']==8
    diagrams();d=Document(TEMPLATE);body=d.element.body;cut=False
    for el in list(body):
        if el.tag==qn('w:p') and el.xpath('./w:pPr/w:pStyle[@w:val="Heading1"]') and ''.join(el.xpath('.//w:t/text()')).startswith('Part B.'):cut=True
        if cut and el.tag!=qn('w:sectPr'):body.remove(el)
    assert cut
    for para in d.paragraphs:
        if para.text.strip()=='PDF Extractor – Project':
            para.runs[0].text='PDF Extractor 项目'
            for run in para.runs[1:]:run.text=''
    d.tables[0].cell(0,1).text='成员英文姓名'
    translations={'PDF Extractor – Project ':'PDF Extractor 项目','Part B.':'Part B.','Functional Prototype':'功能原型','Finished Project (including all coding and materials)':'项目软件与配套材料','Presentation of Project':'项目演示','Students:':'成员信息','Table of Contents':'目录','Student Name (in English)':'成员英文姓名','Student ID No.':'学号'}
    for node in body.xpath('.//w:t'):
        if node.text in translations:node.text=translations[node.text]
    for cell in d.tables[0].rows[0].cells:
        for para in cell.paragraphs:
            for run in para.runs:run.font.color.rgb=RGBColor(255,255,255);run.font.size=Pt(14);run.bold=True
    for section in d.sections:
        for para in section.header.paragraphs:
            para.paragraph_format.line_spacing=1.0
            for run in para.runs:run.font.name='Arial';run.font.size=Pt(10)
    for name,size in [('Normal',11),('Heading 1',18),('Heading 2',15),('Heading 3',12)]:
        st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体' if name=='Normal' else 'Microsoft YaHei')
        st.paragraph_format.space_after=Pt(8)
        if name=='Normal':st.paragraph_format.line_spacing=1.2
        for col in st.element.xpath('.//w:color'):
            for k in ['themeColor','themeShade','themeTint']:col.attrib.pop(qn('w:'+k),None)
    d.core_properties.title='PDF Extractor Part B 功能原型';d.core_properties.language='zh-CN';d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.version='1.0'
    h(d,'Part B. 功能原型',1);h(d,'1. PDF Extractor 项目',1);h(d,'1.1 项目介绍')
    for text in INTRO:p(d,text)
    h(d,'1.2 交付物 2 应用设计',page=True);h(d,'1.2.1 系统架构',3)
    p(d,'工作台采用前后端分离的模块化单体。生产形态由 Nginx 提供 Vite 构建后的静态文件，并将同源 API 转发至 Windows Java 服务；Vite 只参与开发和构建，不承担正式入口。PDF.js 在浏览器渲染预览，服务端根据上传副本生成结果，客户端显示与服务端算法各负其责。')
    figure(d,OUT/'diagrams/architecture.png','图 1 基础 PDF 请求与本机部署边界')
    p(d,'控制器只负责参数绑定与操作分派。PdfService 通过非阻塞信号量限制同时进入的核心请求：默认 2 个，容量已满返回 429。通过校验后创建独立临时目录，检查 PDF 文件头、打开密码和页数，再调用工具类；结果写入响应后执行 finally 清理目录并释放槽位。PDF 解析及输出生成仍位于 Web JVM 内，没有独立 PDF 计算进程，也没有硬执行时间或堆内存隔离。')
    table(d,['组件','实现职责及边界'],[
      ('Windows Web / Worker','Web 监听回环 8080；独立 headless Worker 消费持久后台任务。两类进程可使用不同工作目录。基础 11 项 PDF 操作不入消息队列。'),
      ('PostgreSQL / Flyway','持久记录、任务状态、Outbox 及文件清单；数据库是任务发布与文件就绪状态的事实来源。'),
      ('Redis','共享 HTTP 会话与 Lua 限流；不保存 PDF 原始字节，也不替代数据库记录。'),
      ('RabbitMQ','后台任务使用确认发布、人工 ACK、租约及重试；重复投递由数据库状态与版本约束抑制。'),
      ('Elasticsearch','保存派生全文与向量索引。基础正则检索由 RE2/J 在当前 PDF 文本上完成，不依赖此索引。'),
      ('私有 S3 对象存储','保存需要持久化的文档与内容对象；清单与摘要控制发布和回收。基础即时工具输出通过 HTTP 下载。'),
      ('部署方式','Windows 原生运行 Java，Docker Desktop 的 Linux Engine 提供 Nginx 与支撑容器。服务端口绑定回环；未配置公网 TLS。')],[4,11.9])
    h(d,'1.2.2 类与模块设计',3,page=True)
    p(d,'后端包前缀为 cn.edu.pdfworkbench。模型采用明确的选项与结果类型，控制器和工具库之间由服务层协调。下表中后端相对路径以 backend/src/main/java/cn/edu/pdfworkbench 为根；前端路径以 frontend/src 为根。')
    figure(d,OUT/'diagrams/modules.png','图 2 控制器 服务和 PDF 工具模块')
    module_rows=[(a,b,c.replace('frontend/src/','').replace('；','；\n').replace('/','/\n')) for a,b,c in MODULES]
    table(d,['模块','职责','源码位置'],module_rows,[4.4,6.5,5])
    p(d,'文本处理保留字符位置并聚合连续样式块；阅读顺序由行与栏的空间关系重建。表格处理先收集水平/垂直线段，按连通分量分离网格，再将字符归入单元格。图片处理区别资源身份与绘制次数，原始流和转换结果通过清单关联；转换失败不伪装为成功图像。签名结果将完整性、证书有效期、信任锚和后续修订分列，避免用单一布尔值合并不同判断。')
    h(d,'1.2.3 界面设计 WND 与 Swagger 文档',3,page=True)
    p(d,'界面保留文档预览，将参数与结果置于同一工作区，减少操作切换时重新选择文件。文档页码统一从 1 开始；表格导出索引由接口使用从 0 开始的 tableIndex。工具切换保留文档，新文件替换当前文档；执行期间显示处理中状态，失败后在当前工具修正参数。')
    figure(d,OUT/'diagrams/wnd.png','图 3 与 Part A 对应的 W00 至 W10 导航')
    figure(d,OUT/'evidence/02-text.png','图 4 文本参数与提取结果界面')
    figure(d,OUT/'evidence/03-table.png','图 5 表格与 CSV 导出界面')
    p(d,'Swagger 使用 OpenAPI 3.0.3，核心契约提供 11 个 multipart/form-data POST 操作。同源交互原型位于 /part-a/swagger/，英文提交原型位于 /part-a/submission/swagger/，运行时接口由 /v3/api-docs 提供。接口响应按操作区分 JSON、CSV、ZIP 和 PDF；错误体包含 status、code、message。密码与签名材料以本次请求字段传入，不作为 URL 查询参数或返回字段。')
    spec=json.loads((ROOT/'docs/part-a/submission/openapi.json').read_text('utf-8'))
    rows=[(k.removeprefix('/api/v1/pdf'),', '.join(v['post']['x-requirements']),v['post']['x-wnd']) for k,v in spec['paths'].items()]
    table(d,['POST 后缀','对应需求','WND'],rows,[6.1,7.9,1.9])
    h(d,'1.3 交付物 3 应用软件',page=True);h(d,'1.3.1 可执行软件文件',3)
    p(d,'可执行主体为 Spring Boot JAR 和浏览器静态构建目录，不要求 Windows EXE。运行依赖为 Windows、PowerShell 7.2+、JDK 21+、Maven、Node.js/npm 及已启动的 Docker Desktop Linux Engine。当前编译目标为 Java 21，本轮 JUnit 环境见 1.4.1 节。')
    table(d,['文件或目录','用途'],[
      ('backend/target/pdf-workbench-1.1.0.jar','后端可执行程序，由 Maven package 生成。'),('frontend/dist/','Vue/TypeScript 经 Vite 构建的静态资源，由 Nginx 或 Web 静态处理器提供。'),('deploy/compose.yml','定义本机 Nginx、数据库、缓存、索引、队列及对象存储服务。'),('deploy/Start-Local.ps1','生成本机凭据、启动支撑服务并初始化私有存储桶。'),('deploy/Start-Web.ps1 与 Start-Worker.ps1','启动 Windows 原生 Web 和独立 Worker；读取本机配置。'),('docs/part-a/submission/ 与 docs/part-b/','需求文档、接口契约、设计图及测试证据。')],[7.7,8.2])
    p(d,'从项目根目录依次执行以下构建和启动操作。JAVA_HOME 指向 JDK 21 或更高版本；首次部署需下载依赖和容器镜像。')
    for text in ['npm.cmd --prefix frontend ci','npm.cmd --prefix frontend run build','.\\deploy\\Start-Local.ps1','mvn.cmd -f backend/pom.xml package','.\\deploy\\Start-Web.ps1']:p(d,text)
    p(d,'Web 启动后保持该窗口运行，在另一 PowerShell 会话执行 .\\deploy\\Start-Worker.ps1。浏览器访问 http://127.0.0.1:5173；同机 Vite 开发服务器与 Nginx 不能同时占用 5173。使用 .\\deploy\\Test-Local.ps1 检查依赖服务和入口。')
    p(d,'deploy/.env.local 保存本机生成的配置，不纳入交付包；命名卷保存支撑服务数据。升级 JAR 前先停止对应 Web/Worker 进程，保留数据卷，打包后重新启动。容器删除与数据卷删除不同；数据库和持久对象按同一恢复点配套备份。')
    h(d,'1.4 交付物 4 测试',page=True);h(d,'1.4.1 JUnit 测试及执行记录',3)
    env=evidence['environment']
    p(d,'本轮核心测试于 2026-10-22 执行。JUnit 运行 11 个测试类、93 个实例，失败 0、错误 0、跳过 0；浏览器通过 Nginx 入口运行 8 项工作台用例，全部通过。参数化测试按实例计数。下列 14 条完整记录从本轮执行用例中选取，覆盖正常处理、错误处理和容量边界。')
    table(d,['项目','本轮记录'],[('Java',env.get('java.version','')+' / '+env.get('java.vendor','')),('系统',env.get('os.name','')+' / '+env.get('os.arch','')),('JUnit 汇总','evidence/junit-summary.txt'),('JUnit 用例输出','evidence/core-junit.xml'),('浏览器输出','evidence/browser-summary.txt'),('证据清单','evidence/manifest.json，包含来源报告摘要、时间及源码 SHA-256。')],[4.5,11.4])
    p(d,'JUnit 输出文件保留套件、用例名、执行时间和结果，移除环境属性与控制台输出。JUnit 表中实际结果是已执行断言的结论；PDF 和截图为用例生成的文件。以下 Evidence 路径均以 docs/part-b/evidence 为根；完整用例名可在 core-junit.xml 中检索。')
    for i,c in enumerate(CASES):
        id,title,cls,method,inputs,expect,actual,files=c
        titlep=p(d,id+' '+title);titlep.paragraph_format.keep_with_next=True
        titlep.paragraph_format.page_break_before=(i%3==0)
        for run in titlep.runs:run.bold=True;run.font.size=Pt(12)
        table(d,['记录项','内容'],[('用例',cls+'\n'+method),('输入与步骤',inputs),('预期结果',expect),('实际结果',actual),('状态','通过'),('证据','；'.join(dict.fromkeys(['core-junit.xml']+files.split('；'))))],[2.5,13.4])
    h(d,'1.4.2 可测量验收标准',3,page=True)
    p(d,'验收按输入条件、执行动作与可核对输出定义，需求编号沿用 Part A。范围内条件与算法限制共同组成验收边界。')
    acceptance=[(a,b,c.replace('.searchesWithOffsetsAndRejectsInvalidPatterns','').replace('.reportsPartialImageExportsWithRawFilesAndDiagnostics','').replace('.exposesAllBusinessOperationsInSwagger','')) for a,b,c in ACCEPTANCE]
    table(d,['需求','给定条件与可测量结果','验证定位'],acceptance,[3.0,8.2,4.7])
    p(d,'当前容量上限为 50 MiB 输入、200 页、2 个并行核心请求和 200 MiB 输出；缩小阈值的单元测试验证拒绝与回收逻辑，51 MiB HTTP 上传验证实际上传限制。持续吞吐与最坏响应时间尚无测量基线。')
    h(d,'1.5 交付物 5 源码质量',page=True)
    p(d,'Java 源码使用 Java 21 目标，满足 Java 17 以上的版本要求。控制器、服务、选项/结果模型与 PDF 工具实现分层；try-with-resources 释放文档和流，finally 处理请求目录与配额。输入文件、页码、区域、表达式、口令、证书与输出路径分别校验。修改工具保护源文件，并在失败时避免替换已有输出。')
    mappings=[('FR1 / ERR01','PdfService.run；ApiExceptionHandler','TC09–TC12','core-junit.xml'),('FR2 / FR3 / FR5','PdfTextUtils.extract；TextCollector','TC02','core-junit.xml'),('FR4','PdfTextUtils.search','searchesWithOffsetsAndRejectsInvalidPatterns','core-junit.xml'),('FR6 / FR7 / FR9 / PAGE01','PdfInfoUtils.read','TC14','core-junit.xml'),('FR10 / FR11','PdfTableUtils.extract / writeCsv','TC01、TC03','core-junit.xml；03-table.png'),('FR12','PdfImageUtils.extract / export','TC04；部分转换 HTTP 用例','core-junit.xml；core-partial-images.png'),('FR13–FR16','PdfEncryptionUtils.encrypt','TC05','core-junit.xml'),('FR17–FR19','PdfSignatureUtils.sign / verify / remove','TC06、TC07','signed.pdf；removed.pdf；core-junit.xml'),('WM01 / WM02','PdfWatermarkUtils.addText / addImage','TC08；PdfWatermarkUtilsTest','core-junit.xml'),('NFR1 / NFR4','PdfFiles；PdfService.run','TC05、TC06、TC10–TC12','core-junit.xml'),('WND / 前端','WorkbenchView；useDocument；ResultPanel','8 项 workbench.spec.ts 用例','browser-summary.txt；02-text.png')]
    table(d,['Requirement\n需求','Implementation\n实现','Test Case\n测试用例','Evidence\n证据'],mappings,[2.8,5.2,4.3,3.6])
    p(d,'后端测试源码位于 backend/src/test/java/cn/edu/pdfworkbench，浏览器用例位于 frontend/tests/e2e/workbench.spec.ts。证据清单保存关联源码摘要，便于区分后续修改与本次执行版本。测试夹具在临时目录生成，签名私钥容器不纳入交付；已签 PDF 仅含测试公钥证书与签名结果。')
    h(d,'1.6 交付物 6 范围边界',page=True)
    table(d,['功能','适用范围','限制与排除项'],SCOPE,[2.8,6.5,6.6])
    h(d,'1.7 交付物 7 课堂演示',page=True)
    p(d,'演示按以下 12 页内容组织，围绕同一组样本展示操作、结果及测试证据。现场步骤从工作台文件选择开始，签名材料使用演示专用证书；操作副本与源文件分开保存。')
    table(d,['页','主题','演示内容'],DEMO,[1.0,4.8,10.1])
    p(d,'演示文件使用 api-sample.pdf、signed.pdf、removed.pdf 和 no-images.pdf。正常提取、签名结论、空结果与错误反馈分别展示，避免将业务结果状态与 HTTP 调用失败混为一类。')
    # Chinese run fonts apply only to the authored body and translated cover.
    for run in body.xpath('.//w:r[w:t]'):
        prop=run.get_or_add_rPr();prop.get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei' if run.xpath('ancestor::w:p/w:pPr/w:pStyle[contains(@w:val,"Heading")]') else '宋体')
    OUT.mkdir(exist_ok=True);d.save(OUT/'PartB_Submission.docx')
    print('PartB_Submission.docx generated')
if __name__=='__main__':run()
