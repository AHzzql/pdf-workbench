"""构建 Part A Word、图表和接口原型。"""
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from content import *
import graphics, contract, json

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/part-a'


def set_font(obj,size=12,zh='宋体',bold=False,latin='Times New Roman'):
    obj.font.name=latin;obj.font.size=Pt(size);obj.font.bold=bold;obj.font.color.rgb=RGBColor(0,0,0)
    pr=obj.element.get_or_add_rPr();rf=pr.rFonts
    for k,v in [('ascii',latin),('hAnsi',latin),('eastAsia',zh),('cs',latin)]:rf.set(qn('w:'+k),v)
    for c in pr.findall(qn('w:color')):
        for k in ['themeColor','themeTint','themeShade']:c.attrib.pop(qn('w:'+k),None)
    for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','csTheme']:rf.attrib.pop(qn('w:'+k),None)


def field(p,instruction):
    r=p.add_run();start=OxmlElement('w:fldChar');start.set(qn('w:fldCharType'),'begin');r._r.append(start)
    r=p.add_run();i=OxmlElement('w:instrText');i.set(qn('xml:space'),'preserve');i.text=' '+instruction+' ';r._r.append(i)
    r=p.add_run();sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
    r=p.add_run('1');end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)


def paragraph(d,text,style=None,bold_lead=None):
    p=d.add_paragraph(style=style)
    if bold_lead:
        r=p.add_run(bold_lead);r.bold=True;p.add_run(text)
    else:p.add_run(text)
    return p


def heading(d,text,level=1,page=False):
    p=d.add_heading(text,level);p.paragraph_format.page_break_before=page;return p


def table(d,title,headers,rows,widths):
    p=paragraph(d,title,'Caption');p.paragraph_format.keep_with_next=True
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for c,s in zip(t.rows[0].cells,headers):c.text=str(s)
    for row in rows:
        for c,s in zip(t.add_row().cells,row):c.text=str(s)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'808080');borders.append(x)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
        if ri==0:tr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            mar=OxmlElement('w:tcMar')
            for edge,v in [('top',75),('bottom',75),('left',85),('right',85)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');mar.append(e)
            c._tc.get_or_add_tcPr().append(mar)
            for p in c.paragraphs:
                p.style=d.styles['Table Body'];p.paragraph_format.keep_with_next=(ri==0)
                if ri==0 or ci==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:set_font(r,9.5 if ci==0 and ri>0 and headers[0]=='业务代码' else 10.5,'黑体' if ri==0 else '宋体',ri==0)
    paragraph(d,'','Table Body')
    return t


def figure(d,name,title,width=15.5):
    p=d.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(2)
    p.add_run().add_picture(str(OUT/'diagrams'/name),width=Cm(width))
    paragraph(d,title,'Caption')


def document():
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(2.5);s.bottom_margin=Cm(2.5);s.left_margin=Cm(3);s.right_margin=Cm(2.5);s.header_distance=Cm(1.2);s.footer_distance=Cm(1.2);s.different_first_page_header_footer=True
    for st in d.styles:
        if st.type in [WD_STYLE_TYPE.PARAGRAPH,WD_STYLE_TYPE.CHARACTER]:set_font(st)
        for c in st.element.xpath('.//w:color'):
            c.set(qn('w:val'),'000000')
            for k in ['themeColor','themeTint','themeShade']:c.attrib.pop(qn('w:'+k),None)
        for x in st.element.xpath('.//w:pBdr'):x.getparent().remove(x)
    st=d.styles['Normal'];p=st.paragraph_format;p.first_line_indent=Pt(24);p.line_spacing=1.35;p.space_after=Pt(5);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;p.widow_control=True
    for lvl,size in [(1,16),(2,14),(3,12)]:
        st=d.styles[f'Heading {lvl}'];set_font(st,size,'黑体',True);p=st.paragraph_format;p.first_line_indent=Pt(0);p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.space_before=Pt(12);p.space_after=Pt(7);p.keep_with_next=True;p.keep_together=True
    for n,size in [('Caption',10.5),('Header',9),('Footer',9),('Table Body',10.5),('Small',10.5)]:
        st=d.styles[n] if n in d.styles else d.styles.add_style(n,WD_STYLE_TYPE.PARAGRAPH);set_font(st,size);p=st.paragraph_format;p.first_line_indent=Pt(0);p.line_spacing=1.12;p.space_before=Pt(2);p.space_after=Pt(3);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    d.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for i in [1,2]:
        n=f'TOC {i}';st=d.styles[n] if n in d.styles else d.styles.add_style(n,WD_STYLE_TYPE.PARAGRAPH);set_font(st,11);p=st.paragraph_format;p.first_line_indent=Pt(0);p.left_indent=Pt((i-1)*16);p.line_spacing=1;p.space_before=Pt(0);p.space_after=Pt(0);p.keep_with_next=False
    for n in ['Hyperlink','FollowedHyperlink']:
        if n not in d.styles:d.styles.add_style(n,WD_STYLE_TYPE.CHARACTER)
        set_font(d.styles[n]);d.styles[n].font.underline=False
    set_font(d.styles['Title'],24,'黑体',True);d.styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    h=s.header.paragraphs[0];h.style='Header';h.add_run('PDF 文档处理系统  Part A 需求规格说明书')
    ft=s.footer.paragraphs[0];ft.style='Footer';ft.alignment=WD_ALIGN_PARAGRAPH.CENTER;ft.add_run('第 ');field(ft,'PAGE');ft.add_run(' 页')
    settings=d.settings.element;u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
    cp=d.core_properties;cp.title='PDF 文档处理系统 Part A 需求规格说明书';cp.subject='需求规格、数据流与窗口导航';cp.author='';cp.last_modified_by='';cp.keywords='PDF,Part A,DFD,WND,OpenAPI';cp.comments='';cp.language='zh-CN'
    p=d.add_paragraph();p.paragraph_format.space_before=Pt(65)
    d.add_paragraph('PDF 文档处理系统','Title')
    p=d.add_paragraph('Part A 需求规格说明书');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
    for r in p.runs:set_font(r,18,'黑体')
    for text in ['版本 1.0','2026 年 10 月 20 日']:
        p=d.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(14)
    p=paragraph(d,'本规格定义 PDF 文件选择、内容提取、文档信息、安全处理及水印功能，并建立需求、数据流、界面导航、接口与验收条件之间的对应关系。系统采用 Web 工作台，应用在 Windows 原生运行，数据库与私有对象存储由 Docker 承载。');p.paragraph_format.space_before=Pt(60)
    d.add_page_break();p=d.add_paragraph('目录');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
    for r in p.runs:set_font(r,18,'黑体',True)
    p=d.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);field(p,'TOC \\o "1-2" \\h \\z \\u')
    heading(d,'1 项目概述',page=True)
    heading(d,'1.1 业务问题与目标',2)
    paragraph(d,'PDF 将文字、图像、表格及安全信息组合在页面中。使用者通常需要在多个工具之间切换才能获取内容或生成受保护副本，且提取结果可能丢失页码、位置或表格关系。系统将这些操作集中在一个工作台，保留可核对的页面来源，并将处理结果作为独立文件交付。')
    heading(d,'1.2 系统边界',2)
    paragraph(d,'系统处理单次上传的 PDF 及该操作所需的参数、证书或水印素材，通过浏览器展示信息并下载结果。签署和修改均生成新副本；浏览器的保存设置或保存对话框决定用户下载位置。服务器工作目录和对象存储属于服务内部资源。')
    paragraph(d,'核心范围包括 FR1—FR7、FR9—FR19、WM01、WM02、PAGE01 和 ERR01。FR8 附件提取不纳入当前版本。OCR、任意 PDF 文字编辑、完整 PDF 创作、通用格式转换及在线协作不属于本规格的处理流程。')
    heading(d,'1.3 术语与编号',2)
    table(d,'表 1 术语与编号',['术语','含义'],[
      ('FR','Functional Requirement，沿用任务书的功能需求编号。'),('NFR','Non-Functional Requirement，沿用任务书的非功能需求编号。'),('DFD','Data Flow Diagram，描述数据的输入、处理、存储和输出。'),('WND','Window Navigation Diagram，描述界面与状态之间的导航。'),('WM / PAGE / ERR','水印、总页数与公共错误要求，编号独立于原 FR。'),('默认用户空间','PDF 页面坐标，通常以左下角为原点，单位为 pt；页面旋转与 CropBox 单独传递。'),('空间归属','服务端确定的处理资源边界，不由前端资源标识直接授权。')],[3,12.5])

    heading(d,'2 需求获取与业务约束')
    heading(d,'2.1 利益相关者与使用场景',2)
    table(d,'表 2 利益相关者',['角色','主要需要','关注条件'],STAKEHOLDERS,[3,6,6.5])
    table(d,'表 3 使用场景',['编号','场景','操作与输入','业务输出'],SCENARIOS,[1.6,2.5,6,5.4])
    heading(d,'2.2 需求来源与决策矩阵',2)
    paragraph(d,'功能清单及 FR、NFR 编号以项目任务书为来源。范围选择、部署方式和交叉表含义依据已明确的项目边界；PDF 库、文件格式及执行环境的限制作为技术约束记录。矩阵将宽泛需求转化为输入、输出和可检验条件。')
    table(d,'表 4 需求获取与决策矩阵',['编号','问题','决定或约束','来源与关联'],DECISIONS,[1.3,2.4,8.3,3.5])

    heading(d,'3 功能需求',page=True)
    paragraph(d,'页码表达式使用从 1 开始的文档页码，空表达式表示全部页。文本偏移按 UTF-16 编码单元计数，起点包含、终点不包含。所有操作使用上传副本，输出不得覆盖源文件。')
    for i,r in enumerate(REQUIREMENTS,1):
        heading(d,f'3.{i} {r["id"]} {r["title"]}',2)
        paragraph(d,r['rule'])
        if r['scope']=='不纳入':continue
        paragraph(d,r['input'],bold_lead='输入 ')
        paragraph(d,r['output'],bold_lead='输出 ')
        paragraph(d,r['error'],bold_lead='异常与限制 ')
        paragraph(d,r['accept'],bold_lead=f'{r["case"]} 验收条件 ')

    heading(d,'4 非功能要求与公共错误',page=True)
    heading(d,'4.1 非功能要求',2)
    for id,title,rule,accept in NONFUNCTIONAL:
        heading(d,id+' '+title,3);paragraph(d,rule);paragraph(d,accept,bold_lead='验收安排 ')
    heading(d,'4.2 公共错误目录',2)
    paragraph(d,'错误体由 status、code、message 组成，status 与 HTTP 状态一致。无文本、无表格、无签名和零匹配返回可解释的空结果；图片无法转换使用逐项诊断；验签结论无效并不等同于 HTTP 请求失败。目录中的业务代码用于区分问题，不能用同一条通用消息替代所有类别。')
    table(d,'表 5 公共错误与反馈',['业务代码','状态','触发条件','界面处理'],[(code.replace('CERTIFICATE_PASSWORD_INVALID','CERTIFICATE_PASSWORD_\nINVALID'),status,cause,action) for code,status,cause,action in ERRORS],[5.7,1.7,4.1,4])
    heading(d,'4.3 数据保护与文件生命周期',2)
    paragraph(d,'上传后先建立任务独立暂存区，校验通过后处理。结果对象发布成功后登记归属、类型、长度和校验和，下载经过同一归属校验。失败或超时清理临时输入与未发布输出；清理失败保留可重试记录，不能将其标记为已删除。')
    paragraph(d,'普通工具输入与结果的保存期限默认 24 小时。存在业务引用的持久文件按引用与配额管理；删除先撤销访问，再清理对象和相关记录。数据库与对象存储分别备份，并用对象清单及校验和核对恢复结果。敏感凭据不从历史记录恢复。')

    heading(d,'5 系统架构与技术可行性')
    heading(d,'5.1 模块与部署职责',2)
    paragraph(d,'采用模块化单体。主服务与 PDF 执行入口属于同一代码库，分别承担 HTTP 协调与文档处理；不拆分微服务。Windows 直接运行 Java，Node.js 负责开发和构建；Docker 仅提供 PostgreSQL 与私有对象存储。')
    table(d,'表 6 技术选型与职责',['模块','选型','职责与理由'],ARCHITECTURE,[2.5,4.3,8.7])
    heading(d,'5.2 执行边界与故障处理',2)
    paragraph(d,'主服务使用有界执行器取得处理名额后，以 ProcessBuilder 参数列表启动受限 Java 子进程。PDFBox、签名和图片解码在该进程中执行；密码通过本次受控输入通道传递，不进入命令行。处理期限初值 120 秒，超时终止执行进程并形成明确结果。')
    paragraph(d,'普通操作保持一次请求返回结果的接口语义。有必要后台运行的任务由数据库记录状态，有界执行器仅提供执行名额；领取、租约和提交状态采用短事务，文件读写 不占用数据库行锁。敏感凭据操作在可用执行槽内直接交付；中断后重新输入，不自动重放。')
    paragraph(d,'PostgreSQL 同时承担业务持久化及必要调度，不另设 Redis 或消息队列。SeaweedFS 承担实际文件服务，避免文件长期依赖 Windows 绝对路径。PDF 子进程用于限制内存与故障范围，其运行权限、输入校验和依赖更新分别管理。')

    heading(d,'5.3 本机部署架构',2)
    figure(d,'08-architecture.png','图 1 逻辑与本机部署架构',width=14)
    paragraph(d,'主服务同源提供页面与 API，数据库及 S3 端口只向本机开放。Docker 承载中间件，应用使用 Windows 原生环境。路径、端口和凭据由配置管理，构建产物可迁移至本地虚拟机；公网入口使用 HTTPS。')

    heading(d,'6 数据流设计',page=True)
    heading(d,'6.1 上下文与图例',2)
    paragraph(d,'上下文图把文档使用者作为外部实体，PDF、凭据、操作参数及素材作为输入，将提取结果、处理副本、下载和反馈作为输出。一级图分解系统处理；二级图使用父过程输入与输出端表示分解边界。')
    figure(d,'01-context.png','图 2 系统上下文图')
    table(d,'表 7 数据流图图例',['图形或编号','含义'],[
     ('E1 矩形','文档使用者，输入文件及参数并接收结果。'),('P 编号圆角矩形','把输入数据变换为输出的处理过程。'),('D 编号双线矩形','临时工作区、私有文件及结构化记录。'),('带箭头连线','命名的数据流；状态和错误同样属于输出数据。'),('父过程输入或输出端','二级图与其父过程之间的数据边界，不新增系统外部角色。'),('交叉连线','未标连接点的线条交叉不表示数据合并。')],[4,11.5])
    paragraph(d,'DFD 表示业务数据及其变换。线程、进程和 Docker 的部署关系使用图 1 表达；窗口切换使用第 7 章表达。')
    heading(d,'6.2 一级数据流图',2,page=True)
    paragraph(d,'P1 处理所有输入校验，P2 按操作和范围分发，P3—P5 执行内容提取或文档变换，P6 统一发布结果与状态，P7 授权交付并触发清理。所有失败均经公共结果与错误通道返回。')
    figure(d,'02-level1.png','图 3 一级数据流图')
    paragraph(d,'D1 保存任务期输入和工作数据；D2 保存可授权交付的对象；D3 保存归属、对象元信息及必要执行状态。使用者不直接操作内部存储。',style='Small')
    definitions=[
     ('6.3 文本与文档信息分解','03-text-info.png','图 4 文本与文档信息二级数据流图','P3',[
      ('P3.1','页码、文字、字形几何和样式','规范化文本、块与段落','FR2/FR3/FR5'),('P3.2','规范化页文本与正则选项','匹配文字、偏移和区域','FR4'),('P3.3','目录、元数据和批注结构','页数、信息、书签及批注','FR6/FR7/FR9\nPAGE01')]),
     ('6.4 表格与图片分解','04-tables-images.png','图 5 表格与图片二级数据流图','P4',[
      ('P4.1','页面线段、文字与区域','二维单元格及边界','FR10'),('P4.2','二维单元格和选定索引','保留原行列的呈现及 CSV','FR11'),('P4.3','图片资源与格式选择','原始流、转换图及清单','FR12')]),
     ('6.5 加密与签名分解','05-security.png','图 6 加密与签名二级数据流图','P5',[
      ('P5.1','密码、位数和权限','AES 加密副本','FR13—FR16'),('P5.2','临时签署材料和参数','CMS 签署副本','FR17'),('P5.3','签名内容与可选信任锚','完整性、证书与修订结论','FR18'),('P5.4','未加密 PDF 与确认','移除签名的新副本','FR19')]),
     ('6.6 水印分解','06-watermark.png','图 7 水印二级数据流图','P5',[
      ('P5.5','文字与放置参数','文字水印副本','WM01'),('P5.6','图片与放置参数','等比图片水印副本','WM02')]),
    ]
    for title,image,cap,parent,rows in definitions:
        heading(d,title,2,page=True)
        paragraph(d,f'{parent} 的分解保持父过程接收的 PDF、范围和操作参数，并将对应结果或诊断返回统一交付过程。各操作仅修改任务副本，口令和临时材料不进入持久业务记录。')
        figure(d,image,cap)
        table(d,'表 '+str(8+definitions.index((title,image,cap,parent,rows)))+' 分解过程与需求',['过程','主要输入','主要输出','需求'],rows,[2,5,5.5,3])

    heading(d,'7 窗口导航与界面原型',page=True)
    heading(d,'7.1 导航模型',2)
    paragraph(d,'采用一份当前 PDF 和七组工具面板的工作台。工具切换改变参数与结果视图，保留当前文件。文本提取与正则搜索、签署与验证及移除、文字与图片水印均属于对应工具内的模式切换。')
    figure(d,'07-wnd.png','图 8 窗口与面板导航图')
    paragraph(d,'W08、W09 和 W10 为公共运行状态或面板，不要求打开独立操作系统窗口；其导航与参数面板共同组成 WND。',style='Small')
    heading(d,'7.2 面板与状态定义',2)
    table(d,'表 12 窗口与面板',['编号','面板或状态','输入与动作','反馈'],[
     ('W00','工作台与文件选择','选择、拖入或替换 PDF','文件名、容量、预览和页数'),('W01','文本与搜索','页码、段落选项或表达式','文本、块属性或匹配位置'),('W02','表格','页码、区域及导出索引','二维表格、定位及 CSV'),('W03','图片','PNG/JPG 选择','资源数量、ZIP 和逐项诊断'),('W04','加密','用户与所有者密码、权限、位数','加密副本与明确安全属性'),('W05','数字签名','证书、验证或移除参数','签名副本、分项结论或无签名提示'),('W06','水印','两类素材和放置参数','所选页的新 PDF 副本'),('W07','文档信息','读取当前 PDF','元数据、目录、批注和总页数'),('W08','处理中','防止重复提交，允许离开或取消当前操作','处理状态，释放旧任务订阅'),('W09','结果与下载','检查、定位或保存','实际结果或可解释空集合'),('W10','错误与重新输入','修正页码、口令或替换文件','字段关联错误与可执行操作')],[1.3,3.1,5.4,5.7])
    heading(d,'7.3 文本工具原型',2,page=True)
    figure(d,'09-ui-text.png','图 9 文本工具界面原型')
    table(d,'表 13 文本工具交互',['区域','交互规则'],[
      ('当前文档','更换文件后清除旧结果和口令；显示新文件名与可读取页数。'),('参数','提取与搜索共用页码；切换模式时更新参数并清除不适用结果。'),('文本块属性','逐块展示文字、字体、字号、颜色与 X/Y/宽/高，选择后定位页面。'),('结果','页面文字保留换行及段落；空白页显示无可提取文本；下载使用浏览器保存流程。')],[3.2,12.3])
    heading(d,'7.4 数字签名原型',2,page=True)
    figure(d,'09-ui-signature.png','图 10 数字签名验证界面原型')
    table(d,'表 14 验签展示',['维度','显示规则'],[
      ('完整性','分别显示有效、无效或无法验证；无法验证说明原因。'),('后续修订','区分存在、不存在和无法判定，不用 false 代替未知。'),('证书','主体、签发者、序列号和有效时间可展开；有效期和信任独立展示。'),('移除确认','说明会改变副本并取消签名效力，明确确认后提交，原文件保留。')],[3.2,12.3])

    heading(d,'8 接口与 Swagger 原型',page=True)
    heading(d,'8.1 核心接口目录',2)
    paragraph(d,'核心前缀为 /api/v1/pdf。以下操作均使用 POST 和 multipart/form-data；file 为 PDF 必填项，password 为可选打开密码。业务参数放入正文而非 URL 查询串。成功时返回 JSON 或带附件文件名的二进制结果。')
    table(d,'表 15 核心接口',['编号','相对路径','功能','成功类型'],[(i,p,t,'JSON' if k not in ['pdf','csv','zip'] else k.upper()) for i,p,t,r,w,o,k in INTERFACES],[1.8,4.7,6.7,2.3])
    heading(d,'8.2 公共字段与条件',2)
    table(d,'表 16 关键参数约定',['操作','参数','规则'],[
      ('全部','file / password','PDF 必填，打开密码可选；密码字段 writeOnly。'),('文字与表格','pages','空值为全部页；页码从 1 开始；逗号和闭区间允许组合。'),('搜索','pattern / flags','表达式非空；支持 RE2/J 及 1、2、4 的按位组合。'),('表格','x / y / width / height','区域参数同时提供；tableIndex 从 0 开始。'),('加密','ownerPassword / userPassword / keyLength','所有者密码必填且不同；位数为 128 或 256；权限为三个布尔值。'),('签署','certificate / certificatePassword / alias','PKCS#12 与口令必填，多私钥条目需明确别名。'),('验证','trustCertificates','可选 X.509 PEM/DER；未提供时信任为未评估。'),('移除','confirmRemoval','必须为 true，且输入满足未加密和修改权限限制。'),('水印','type 及素材','TEXT 必填 text；IMAGE 必填 image；共用 pages、位置、角度及透明度。')],[2.2,5,8.3])
    heading(d,'8.3 结果类型与错误契约',2)
    table(d,'表 17 关键结果模型',['模型','内容'],[
      ('TextResult / PageText','分页文字、页面几何、spans、blocks、paragraphs 与 diagnostics。'),('TextBlock','page、text、UTF-16 start/end、bounds、font、fontSize 和 fillColor。'),('SearchMatch','页码、匹配文本、偏移及矩形集合。'),('DocumentInfo','pageCount、标准/自定义元数据、XMP、bookmarks、annotations 和诊断。'),('TableData','page、bounds 与 rows 二维数组；无表返回 []。'),('ImageManifest','图片资源信息、occurrences、rawFile、转换 file、EXPORTED/RAW_ONLY 及诊断。'),('SignatureResult','完整性、可空后续修订、证书有效期/信任、证书摘要和原因。'),('Error','status、code、message，媒体类型为 application/json。')],[4.2,11.3])
    paragraph(d,'OpenAPI 3.0.3 契约位于 openapi.json；swagger/index.html 使用同一份模型生成接口界面。Swagger 展示请求与响应契约，WND 表达窗口导航，两者通过 API 和 W 编号关联。ZIP 内清单模型与图片总数、成功转换数响应头共同表达部分成功情况。')
    paragraph(d,'字段约束和条件以第 3 章及 OpenAPI 描述共同定义；XML、PDF 结构和证书由处理模块验证。页面的输入限制不能代替服务端校验。',style='Small')

    heading(d,'9 验收与需求追踪')
    heading(d,'9.1 验收方法',2)
    paragraph(d,'功能验收使用文字、位置、表格矩阵和图片像素已知的合成 PDF。预期值由样本参数和独立读取方式确定。签名验证分别使用正常签名、签名字节篡改、签后追加修订和不同证书条件；修改输出重新打开并检查内容、权限或签名结构。')
    paragraph(d,'第 3 章的 AC01—AC22 为功能验收条件。执行层安排 JUnit 单元及 HTTP 契约用例，界面层安排上传、参数、定位和下载用例。性能、资源故障与备份恢复按 NFR 独立记录环境、样本、预期与实际结果。')
    heading(d,'9.2 功能追踪矩阵',2)
    table(d,'表 18 功能追踪',['需求','DFD','WND','接口','验收'],[(r['id'],r['dfd'],r['wnd'],r['api'],r['case']) for r in REQUIREMENTS],[1.8,2.9,3.3,5.5,2])
    heading(d,'9.3 非功能追踪',2)
    table(d,'表 19 非功能追踪',['需求','关联要求','验证安排'],[
     ('NFR1','FR1、全部结果与 W00—W10','分辨率、键盘导航、错误关联和真实下载'),('NFR2','A03、A18、FR2/12','20 页固定样本、容量与并发边界'),('NFR3','ERR01、P1/P6/P7','损坏、空白、子进程异常、超时与清理'),('NFR4','FR13—19、WM01/02','源文件校验和、归属、敏感材料与日志'),('NFR5','图 1、表 6 与核心契约','模块依赖、独立 JUnit、契约与编号一致性'),('NFR6','A01、A17、配置与依赖','Windows 启动、Docker 连接及虚拟机迁移检查')],[2,5,8.5])
    heading(d,'10 参考资料',page=True)
    for id,title,value in REFERENCES:
        p=paragraph(d,id+' '+title);p.paragraph_format.keep_with_next=True
        paragraph(d,value,'Small')
    d.save(OUT/'PartA需求规格说明书.docx')
    return d


def main():
    OUT.mkdir(parents=True,exist_ok=True);figs=graphics.build(OUT/'diagrams');spec=contract.build(OUT);document()
    links=''.join(f'<li><a href="diagrams/{f["id"]}.svg">{f["title"]}</a></li>' for f in figs)
    (OUT/'index.html').write_text('''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>PDF 文档处理 Part A</title><link rel="icon" href="data:,"><style>body{max-width:920px;margin:52px auto;padding:0 24px;font:17px/1.9 "Microsoft YaHei",sans-serif;color:#222}h1{font-size:30px}h2{margin-top:32px;font-size:22px}a{color:#222;text-underline-offset:4px}li{margin:8px 0}footer{border-top:1px solid #ccc;margin-top:40px;padding-top:15px;font-size:14px}</style><h1>PDF 文档处理系统</h1><p>Part A 需求规格说明书与接口设计</p><h2>文档与接口</h2><ul><li><a href="PartA需求规格说明书.docx">Part A 需求规格说明书</a></li><li><a href="swagger/index.html">Swagger 核心接口原型</a></li><li><a href="openapi.json">OpenAPI 3.0.3 契约</a></li></ul><h2>数据流与界面图</h2><ul>'''+links+'''</ul><footer>版本 1.0　2026 年 10 月 20 日</footer></html>''',encoding='utf-8')
    (OUT/'README.md').write_text('''# PDF 文档处理系统 Part A\n\n`PartA需求规格说明书.docx` 为主文档；`openapi.json` 为核心接口设计契约；`swagger/index.html` 为本地 Swagger 原型；`diagrams` 保存 SVG 图源、PNG 及节点连线数据。\n\n可直接打开 `index.html` 查看交付目录。Swagger 使用本地资源，不依赖 CDN；执行请求时连接契约中配置的本机应用地址。\n\n编辑需求与图形数据后，使用带 python-docx 和 Pillow 的 Python 运行 `tools/part_a/build.py` 重建，随后更新 Word 目录并重新渲染检查。\n''',encoding='utf-8')
    print(json.dumps({'document':str(OUT/'PartA需求规格说明书.docx'),'requirements':len(REQUIREMENTS),'figures':len(figs),'operations':len(spec['paths'])},ensure_ascii=False))

if __name__=='__main__':main()
