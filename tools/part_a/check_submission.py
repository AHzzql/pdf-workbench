"""Check the English submission against its template and interface baseline."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from lxml import etree as E
from PIL import Image
import json,re,hashlib
from submission_content import PROJECT_INTRO,FUNCTIONAL_INTRO,AMBIGUITIES,SOFTWARE,ERRORS,GROUPS,NEEDS

ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'docs/part-a/submission'
QA=ROOT/'backend/target/part-a-submission-qa'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
def run():
    checks=[]
    def check(name,condition,detail=''):
        checks.append({'check':name,'passed':bool(condition),'detail':detail})
    source=ROOT/'docs/templates/PartA-Submission-Template.docx';path=OUT/'PartA_Submission.docx'
    template=Document(source);doc=Document(path)
    norm=lambda v:re.sub('[^a-z0-9]','',v.lower())
    headings=[p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
    expected=[p.text for p in template.paragraphs if p.style.name.startswith('Heading')]
    position=-1
    for title in expected:
        indices=[i for i,t in enumerate(headings) if norm(t)==norm(title) and i>position]
        check('Template heading: '+title.strip(),bool(indices))
        if indices:position=indices[0]
    for part in ('word/document.xml','word/header1.xml','word/footer1.xml'):
        with ZipFile(path) as z:
            root=E.fromstring(z.read(part));text=' '.join(root.xpath('//w:t/text()',namespaces=NS))
            check('English text: '+part,not re.search('[\u3400-\u9fff]',text))
            if part.endswith('document.xml'):
                check('TOC field',bool(root.xpath('//w:instrText[contains(text(),"TOC")]',namespaces=NS)))
                check('No field errors','Error! Bookmark' not in text and 'Error! Reference' not in text)
                check('No unfilled instructional prose',not re.search(r'\bhere you should\b|\bStudents should\b|\bas requested\b|\bAI\b|\bassistant\b',text,re.I))
                check('Ten embedded figures',len(root.xpath('//w:drawing',namespaces=NS))==10)
                check('A4 page',root.xpath('//w:pgSz/@w:w',namespaces=NS)==['11906'] and root.xpath('//w:pgSz/@w:h',namespaces=NS)==['16838'])
    check('English document metadata',doc.core_properties.language=='en-GB')
    check('Original member table structure',len(doc.tables[0].rows)==6 and len(doc.tables[0].columns)==3)
    check('No invented member identities',all(not cell.text for row in doc.tables[0].rows[1:] for cell in row.cells[1:]))
    all_text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
    for identifier,paragraphs in [('1.1',PROJECT_INTRO),('2.1',FUNCTIONAL_INTRO)]:
        words=len(re.findall(r"\b[\w]+(?:[-’'][\w]+)*\b",' '.join(paragraphs)))
        check(identifier+' introduction word count',200<=words<=300,str(words))
        check(identifier+' introduction present',all(t in all_text for t in paragraphs))
    for label,rows in [('A',AMBIGUITIES),('Q',SOFTWARE)]:
        matching=[t for t in doc.tables if len(t.rows)>1 and t.rows[1].cells[0].text==rows[0][0]]
        check(label+' independent assumption grid',len(matching)==1)
        if matching:
            check(label+' grid covers all questions',[r.cells[0].text for r in matching[0].rows[1:]]==[r[0] for r in rows])
    for i,text in enumerate(NEEDS,1):check('User need '+str(i),text in all_text)
    for condition,code,message in ERRORS:check('Error package: '+condition,all(t in all_text for t in [condition,code,message]))
    spec=json.loads((OUT/'openapi.json').read_text('utf-8'))
    old=json.loads((ROOT/'docs/part-a/openapi.json').read_text('utf-8'))
    def shape(v):
        if isinstance(v,dict):return {k:shape(x) for k,x in v.items() if k not in {'title','description','summary','tags','info','x-wnd'}}
        if isinstance(v,list):return [shape(x) for x in v]
        return v
    check('Core API constraints unchanged',shape(spec)==shape(old))
    check('English OpenAPI',not re.search('[\u3400-\u9fff]',json.dumps(spec,ensure_ascii=False)))
    check('Eleven POST operations',len(spec['paths'])==11 and all(set(o)=={'post'} for o in spec['paths'].values()))
    check('Same-origin gateway',spec['servers'][0]['url']=='/')
    def walk(v):
        if isinstance(v,dict):
            if '$ref' in v:
                target=spec
                for key in v['$ref'][2:].split('/'):target=target[key]
                check('Schema reference '+v['$ref'],bool(target))
            for key,x in v.items():
                if key=='password' or key.endswith('Password'):
                    if isinstance(x,dict):check('Write-only credential '+key,x.get('writeOnly') is True and 'default' not in x and 'example' not in x)
                walk(x)
        elif isinstance(v,list):
            for x in v:walk(x)
    walk(spec)
    for route,operation in spec['paths'].items():
        suffix=route.removeprefix('/api/v1/pdf')
        panel='W02' if suffix.startswith('/text') else 'W03' if suffix.startswith('/tables') else 'W04' if suffix=='/images' else 'W05' if suffix=='/info' else 'W06' if suffix=='/encrypt' else 'W07' if suffix.startswith('/signatures') else 'W08'
        check('WND link '+route,operation['post']['x-wnd']==panel)
    script=(OUT/'swagger/contract.js').read_text('utf-8').split('=',1)[1].strip().removesuffix(';')
    check('Swagger uses English contract',json.loads(script)==spec)
    html=(OUT/'swagger/index.html').read_text('utf-8')
    check('Offline Swagger assets',all((OUT/'swagger'/s).resolve().is_file() for s in re.findall(r'(?:src|href)="([^"]+\.(?:css|js))"',html)))
    for group in GROUPS:
        name=group[1];data=json.loads((OUT/'diagrams'/f'{name}.json').read_text('utf-8'))
        check('English DFD '+name,not re.search('[\u3400-\u9fff]',json.dumps(data,ensure_ascii=False)))
        ids={n['id'] for n in data['nodes']}
        check('DFD edge endpoints '+name,all(e['from'] in ids and e['to'] in ids for e in data['edges']))
        check('DFD named flows '+name,all(e['label'] for e in data['edges']))
        check('DFD process identities '+name,len(ids)==len(data['nodes']) and all(n['label'].strip() and not re.match(r'^(?:P\d|E\d|W\d)',n['label']) for n in data['nodes']))
        check('DFD included in document '+name,group[0] in all_text)
        check('Temporary input store '+name,any(n['id']=='D1' and 'Temporary input copy' in n['label'] for n in data['nodes']))
    wnd=json.loads((OUT/'diagrams/wnd.json').read_text('utf-8'))
    check('All tool panels in WND',set('W'+str(i).zfill(2) for i in range(11))=={n['id'] for n in wnd['nodes']})
    check('Retry navigation',any(e['from']=='W10' and e['to']=='W01' for e in wnd['edges']))
    check('Swagger evidence image',(OUT/'diagrams/swagger-prototype.png').is_file())
    image=Image.open(OUT/'diagrams/swagger-prototype.png');check('Swagger capture dimensions',image.size==(1280,1000))
    pages=json.loads((QA/'page-index.json').read_text('utf-8'))
    render=json.loads((QA/'render.json').read_text('utf-8-sig'))
    raster=json.loads((QA/'raster.json').read_text('utf-8'))
    pdf_hash=hashlib.sha256((QA/'PartA.pdf').read_bytes()).hexdigest()
    check('Render uses current Word and PDF',render['WordSHA256'].lower()==hashlib.sha256(path.read_bytes()).hexdigest() and render['PdfSHA256'].lower()==pdf_hash)
    check('Raster uses current PDF',raster['pdf_sha256']==pdf_hash)
    check('Rendered page counts agree',render['Pages']==raster['pages']==len(pages))
    check('All rendered pages available',all((QA/'pages'/f'page-{p["page"]}.png').is_file() for p in pages))
    check('No sparse accidental page',all(p['characters']>=200 for p in pages),str(len(pages)))
    result={'checks':len(checks),'passed':sum(c['passed'] for c in checks),'failures':[c for c in checks if not c['passed']],
            'pages':len(pages),'word_sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
    QA.mkdir(parents=True,exist_ok=True);(QA/'checks.json').write_text(json.dumps({'summary':result,'checks':checks},indent=2),encoding='utf-8')
    print(json.dumps(result));return not result['failures']
if __name__=='__main__':raise SystemExit(0 if run() else 1)
