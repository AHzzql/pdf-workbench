"""Build the English Part A using the supplied submission template."""
from pathlib import Path
import json,re,hashlib
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
import graphics
from submission_content import *

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/part-a/submission'
TEMPLATE=ROOT/'docs/templates/PartA-Submission-Template.docx'
CJK=re.compile(r'[\u3400-\u9fff]')

class EnglishFigure(graphics.Figure):
    def node(self,id,text,x,y,w=350,h=110,kind='process'):
        text=re.sub(r'^(?:W\d+|E\d+[a-z]?|P\d+(?:\.\d+)*|D\d+)\s+','',text)
        return super().node(id,text,x,y,w,h,kind)
    def text(self,x,y,value,size=30,anchor='mm',background=False):
        super().text(x,y,value,max(24,min(size,28)),anchor,background)
    def save(self,out):
        result=super().save(out)
        p=out/(self.name+'.svg')
        p.write_text(p.read_text('utf-8').replace('SimHei,Microsoft YaHei,sans-serif','Arial,Helvetica,sans-serif'),encoding='utf-8')
        return result

def diagrams():
    target=OUT/'diagrams';target.mkdir(parents=True,exist_ok=True)
    graphics.FONT='C:/Windows/Fonts/arial.ttf'
    result=[]
    f=EnglishFigure('context','PDF Extractor context',1400,620)
    f.node('E1','E1 Document user',230,290,370,150,'external')
    f.node('P0','P0 PDF Extractor',1010,290,560,200)
    f.edge('E1','P0','PDF, credentials and parameters',[(415,240),(730,240)],(570,170),23)
    f.edge('P0','E1','Results, new files and feedback',[(730,340),(415,340)],(570,420),23)
    f.text(700,550,'Operations use copies; the source document is not overwritten.',24)
    result.append(f.save(target))
    for i,g in enumerate(GROUPS,2):
        title,name,intro,branches,*_=g
        f=EnglishFigure(name,title,1400,1120)
        f.node('E1','E1 Document user',700,75,430,100,'external')
        f.node('P1','P1 Validate input\nFR1 and common constraints',700,270,430,120)
        f.node('P9','P9 Explain errors\nCommon package',1140,270,360,120)
        f.node('D1','Temporary input copy\nOne operation only',250,270,350,120,'store')
        f.edge('P1','D1','Staged copy',[(485,245),(425,245)],(310,185),21)
        f.edge('D1','P1','Input bytes',[(425,295),(485,295)],(280,355),21)
        f.edge('E1','P1','PDF and operation inputs',[(700,125),(700,210)],(700,165),22)
        f.edge('P1','P9','Invalid input',[(915,270),(960,270)],(1095,180),22)
        sequential=name not in {'text','signatures'}
        process_ids=[value.split()[0] for value in branches]
        for j,(x,text) in enumerate(zip([250,700,1150],branches)):
            ident=process_ids[j];f.node(ident,text,x,525,400,125)
            pts=[(700,330),(700,400),(x,400),(x,462)] if x!=700 else [(700,330),(700,462)]
            if not sequential or j==0:f.edge('P1',ident,'PDF and parameters',pts,(x,435),21)
        if sequential:
            flows={'tables':('Line segments','Cell grid'),'images':('Resources','Conversions'),
                   'encryption':('Passwords','Permissions'),'watermark':('Selection','Mark settings')}[name]
            f.edge(process_ids[0],process_ids[1],flows[0],[(450,525),(500,525)],(475,430),20)
            f.edge(process_ids[1],process_ids[2],flows[1],[(900,525),(950,525)],(925,430),20)
        f.text(700,380,'Validated input copy and selected parameters',22,background=True)
        f.node('P8','P8 Assemble result\nand download metadata',700,800,440,120)
        for j,x in enumerate([250,700,1150]):
            if not sequential or j==2 or (name=='tables' and j==1):
                label={'text':['Page text','Pattern matches','PDF properties'],
                       'signatures':['Signed copy','Verification records','Unsigned copy'],
                       'tables':['','Cell rows','CSV output'],'images':['','','Image package'],
                       'encryption':['','','Encrypted copy'],'watermark':['','','Marked copy']}[name][j]
                f.edge(process_ids[j],'P8',label,[(x,588),(x,675),(700,675),(700,740)] if x!=700 else [(700,588),(700,740)],(x,620),21)
        f.text(700,665,'Structured data or a processed copy',22,background=True)
        f.node('E1b','E1 Document user',700,1020,470,100,'external')
        f.edge('P8','E1b','Result and status',[(700,860),(700,970)],(700,915),22)
        f.edge('P9','E1b','Error details',[(1320,270),(1365,270),(1365,1020),(935,1020)],(1240,925),22)
        result.append(f.save(target))
    f=EnglishFigure('wnd','Window and panel navigation',1400,1390)
    f.node('W00','W00 Select / open PDF',700,70,460,95,'external')
    f.node('W01','W01 Preview and page navigation\nPDF loaded; page total displayed',700,260,540,130)
    f.edge('W00','W01','Valid selected file',[(700,118),(700,195)],(700,153),22)
    panels=['W02 Text and search','W03 Tables and CSV','W04 Images and ZIP','W05 Document information','W06 Encryption / permissions','W07 Digital signatures','W08 Text / image watermark']
    for j,label in enumerate(panels):
        y=465+j*120;ident='W'+str(j+2).zfill(2)
        f.node(ident,label,405,y,470,80)
        f.edge('W01',ident,'',[(430,260),(90,260),(90,y),(170,y)])
        f.edge(ident,'W09','',[(640,y),(770,y),(770,750),(830,750)])
    f.node('W09','W09 Inspect results\nSave supported output',1060,750,460,140)
    f.node('W10','W10 Read status / errors\nCorrect inputs and retry',1060,1110,460,140)
    f.edge('W09','W10','Failure feedback',[(1060,820),(1060,1040)],(1060,920),22)
    f.edge('W10','W01','Correct and retry',[(1060,1180),(1060,1250),(40,1250),(40,260),(430,260)],(950,1250),22)
    f.edge('W09','W01','',[(1290,750),(1340,750),(1340,260),(970,260)])
    f.text(1110,190,'Choose another tool',22,background=True)
    f.text(700,1340,'Tool changes preserve the selected document; a new upload replaces it.',23)
    result.append(f.save(target))
    f=EnglishFigure('text-prototype','Text tool wireframe',1400,1000)
    f.text(700,55,'PDF Extractor workbench',30)
    f.node('file','Selected PDF: sample.pdf     Page 1 / 3     Choose PDF',700,140,1300,90,'external')
    f.node('tools','Text | Tables | Images | Info | Security | Signatures | Watermark',700,245,1300,85,'external')
    f.node('preview','PDF page preview\n\nAlpha is 42.\n\nSource page coordinates\nremain available for results.',370,600,600,570,'external')
    f.node('params','Text extraction\nPages: 1,3\nPreserve paragraphs: yes',1035,410,600,190,'external')
    f.node('action','Extract text     Search pattern',1035,570,600,85,'external')
    f.node('result','Page 1: Alpha is 42.\nBlocks / bounds / style\nStatus: extraction completed',1035,745,600,225,'external')
    f.text(700,945,'Conceptual interface prototype using the same panels as the navigation diagram.',22)
    result.append(f.save(target))
    (OUT/'figures.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    return result

TRANSLATIONS={
'1 忽略大小写，2 点匹配换行，4 多行；支持按位组合。':'1: ignore case; 2: dot matches line breaks; 4: multiline. Combine flags bitwise.',
'PDF 文档处理核心接口设计':'PDF Extractor Core Interface Design',
'PDF 默认用户空间；四项必须同时提供。':'PDF default user space; supply all four values together.',
'PDF，最大 50 MiB、200 页。':'PDF, at most 50 MiB and 200 pages.',
'RE2/J 支持的正则表达式。':'A regular expression supported by RE2/J.',
'VALID 有效；INVALID 无效；UNKNOWN 无法验证。':'VALID: verified integrity; INVALID: invalid integrity; UNKNOWN: indeterminate verification.',
'ZIP 包包含 manifest.json、原始流及可转换的图片；清单模型见 ImageManifest。':'ZIP containing manifest.json, original streams and available image conversions. See ImageManifest.',
'attachment 及安全文件名。':'attachment and a safe output filename.',
'type=IMAGE 时必填，PNG/JPG，最大 20 MiB、4000 万像素。':'Required for type=IMAGE; PNG/JPG, at most 20 MiB and 40 million pixels.',
'内容提取':'Content extraction',
'分页文本、表格、图片、文档信息、加密、签名与水印操作。接口字段与 Part A 需求编号关联。':'Page text, tables, images, document information, encryption, signatures and watermarks. Parameters link to Part A requirement IDs.',
'可选 X.509 PEM/DER 信任证书，最大 5 MiB。':'Optional X.509 PEM/DER trust certificates, at most 5 MiB.',
'图片资源数量。':'Number of image resources.',
'处理成功。':'Operation succeeded.',
'处理错误，具体类别见 code 和 message。':'Operation error; see code and message for the specific category.',
'安全与修改':'Security and modification',
'当前同源入口':'Current same-origin gateway',
'成功转换的图片资源数量。':'Number of successfully converted image resources.',
'所有参数均在 multipart 正文中传递。':'Supply all parameters in the multipart request body.',
'所有参数均在 multipart 正文中传递。TEXT 必填 text；IMAGE 必填 image；两类水印共同使用页码、位置、角度和透明度。':'Supply all parameters in the multipart body. TEXT requires text; IMAGE requires image. Both use page selection, position, angle and opacity.',
'文档信息':'Document information',
'无法判定时为 null。':'Null when the result cannot be determined.',
'明确确认移除副本中全部签名。':'Explicit confirmation to remove all signatures from the copy.',
'最大 5 MiB 的 PKCS#12 容器。':'PKCS#12 container, at most 5 MiB.',
'本次操作使用的打开密码。':'Opening password for this operation only.',
'留空表示无需密码打开。':'An empty value permits opening without a password.',
'空值表示全部页；文档页码从 1 开始，例如 1,3-5。':'Empty selects all pages. Document page numbers start at 1, for example 1,3-5.',
'非空且必须不同于用户密码。':'Non-empty and different from the user password.'}

def contract():
    spec=json.loads((ROOT/'docs/part-a/openapi.json').read_text('utf-8'))
    def translate(value):
        if isinstance(value,dict):return {k:translate(v) for k,v in value.items()}
        if isinstance(value,list):return [translate(v) for v in value]
        if isinstance(value,str) and CJK.search(value):
            m=re.match(r'API(\d+) ',value)
            if m:return 'API'+m[1]+' '+API_NAMES[int(m[1])-1]
            return TRANSLATIONS[value]
        return value
    spec=translate(spec)
    for path,item in spec['paths'].items():
        suffix=path.removeprefix('/api/v1/pdf')
        item['post']['x-wnd']='W02' if suffix.startswith('/text') else 'W03' if suffix.startswith('/tables') else 'W04' if suffix=='/images' else 'W05' if suffix=='/info' else 'W06' if suffix=='/encrypt' else 'W07' if suffix.startswith('/signatures') else 'W08'
    (OUT/'openapi.json').write_text(json.dumps(spec,indent=2),encoding='utf-8')
    sw=OUT/'swagger';sw.mkdir(exist_ok=True)
    (sw/'contract.js').write_text('window.PDF_CORE_OPENAPI = '+json.dumps(spec)+';\n',encoding='utf-8')
    (sw/'index.html').write_text('''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>PDF Extractor Core API</title><link rel="icon" href="data:,"><link rel="stylesheet" href="../../swagger/vendor/swagger-ui.css"><style>body{margin:0}header{padding:16px 28px;font:16px Arial;border-bottom:1px solid #bbb}header a{margin-right:28px;color:#222}.swagger-ui .topbar{display:none}.swagger-ui .opblock{box-shadow:none}.swagger-ui .opblock.opblock-post{border-color:#555;background:#fafafa}.swagger-ui .opblock.opblock-post .opblock-summary-method{background:#333}</style></head><body><header><a href="../index.html">Part A submission</a><a href="../openapi.json">OpenAPI JSON</a></header><main id="swagger-ui"></main><script src="../../swagger/vendor/swagger-ui-bundle.js"></script><script src="contract.js"></script><script>window.ui=SwaggerUIBundle({spec:window.PDF_CORE_OPENAPI,dom_id:'#swagger-ui',deepLinking:true,docExpansion:'list',defaultModelsExpandDepth:0,validatorUrl:null,persistAuthorization:false});</script></body></html>''',encoding='utf-8')
    return spec

def p(d,text,style=None):
    return d.add_paragraph(text,style)
def h(d,text,level=1,page=False):
    para=d.add_heading(text,level);para.paragraph_format.page_break_before=page;return para
def table(d,headers,rows,widths,compact=False):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    grid=t._tbl.tblGrid
    for child,w in zip(grid,widths):child.set(qn('w:w'),str(round(w*567)))
    for cell,text in zip(t.rows[0].cells,headers):cell.text=text
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    border=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');border.append(b)
    t._tbl.tblPr.append(border)
    for i,row in enumerate(t.rows):
        if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,cell in enumerate(row.cells):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1 if compact else 3);para.paragraph_format.space_before=Pt(1 if compact else 3)
                para.paragraph_format.line_spacing=1.05
                for run in para.runs:run.font.name='Arial';run.font.size=Pt(9.5 if compact and j==1 and i else 10 if compact else 10.5);run.font.color.rgb=RGBColor(0,0,0);run.bold=(i==0)
    p(d,'')
def fig(d,name,caption):
    para=p(d,'');para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next=True
    para.add_run().add_picture(str(OUT/'diagrams'/f'{name}.png'),width=Cm(16))
    cap=p(d,caption);cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.space_after=Pt(12)
    for run in cap.runs:run.italic=True;run.font.size=Pt(10.5)

def document(spec):
    d=Document(TEMPLATE)
    body=d._element.body;cut=False
    for el in list(body):
        if el.tag==qn('w:p'):
            text=''.join(el.xpath('.//w:t/text()'))
            if text.startswith('Part A.') and 'Requirements Specification' in text and el.xpath('./w:pPr/w:pStyle[@w:val="Heading1"]'):
                cut=True
        if cut and el.tag!=qn('w:sectPr'):body.remove(el)
    # The body-level text test also covers split runs through the OOXML text nodes.
    if not cut:raise AssertionError('Template main body boundary not found')
    for cell in d.tables[0].rows[0].cells:
        for para in cell.paragraphs:
            for run in para.runs:run.font.color.rgb=RGBColor(255,255,255)
    for section in d.sections:
        for para in section.header.paragraphs:
            para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.0
            for run in para.runs:run.font.name='Arial';run.font.size=Pt(10)
    if 'List Bullet' not in d.styles:
        st=d.styles.add_style('List Bullet',WD_STYLE_TYPE.PARAGRAPH);st.base_style=d.styles['Normal']
        numbering=d.part.numbering_part.element
        abstract=numbering.xpath('./w:abstractNum[w:lvl/w:numFmt[@w:val="bullet"]]/@w:abstractNumId')[0]
        num=numbering.xpath(f'./w:num[w:abstractNumId[@w:val="{abstract}"]]/@w:numId')[0]
        prop=st.element.get_or_add_pPr();np=OxmlElement('w:numPr')
        level=OxmlElement('w:ilvl');level.set(qn('w:val'),'0');np.append(level)
        ident=OxmlElement('w:numId');ident.set(qn('w:val'),num);np.append(ident);prop.append(np)
    for name,size in [('Normal',11),('Heading 1',18),('Heading 2',15),('Heading 3',12)]:
        st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(8)
        if name=='Normal':st.paragraph_format.line_spacing=1.12
        for col in st.element.xpath('.//w:color'):
            for key in ['themeColor','themeShade','themeTint']:col.attrib.pop(qn('w:'+key),None)
    d.core_properties.title='PDF Extractor Part A Requirements Specification and Design'
    d.core_properties.language='en-GB';d.core_properties.version='1.2'
    h(d,'Part A. Requirements Specification / Design:')
    h(d,'1. PDF Extractor Project')
    h(d,'1.1. Introduction',2)
    for text in PROJECT_INTRO:p(d,text)
    h(d,'1.2 Key Feature List',2)
    for text in FEATURES:p(d,text,'List Bullet')
    h(d,'1.3 Requirements Elicitation Work',2)
    h(d,'1.3.1 Stakeholders',3)
    table(d,['Stakeholder','Example role','Required outcome'],STAKEHOLDERS,[3.6,5.2,7.1])
    h(d,'1.3.2 User Needs Identified',3)
    for text in NEEDS:p(d,text,'List Bullet')
    h(d,'1.3.3 Requirements Questions',3)
    p(d,'The feature statements leave choices about representation, duplicates, empty results and security conclusions. The following grid records the selected interpretation and distinguishes scope decisions from technical constraints.')
    types=lambda v:'Decision' if 'decision' in v.lower() else 'Scope' if 'scope' in v.lower() else 'Need' if ('need' in v.lower() or 'requirement' in v.lower()) else 'Criterion' if ('acceptance' in v.lower() or 'result' in v.lower()) else 'Constraint'
    table(d,['ID','Question / ambiguity','Recorded resolution','Type','Source'],[(a,b,c,types(t),s) for a,b,c,t,s in AMBIGUITIES],[1.2,3.2,6.7,2.2,2.6])
    h(d,'1.3.4 Software Requirements Elicitation Questions',3)
    p(d,'This grid covers the operating context and acceptance boundaries. Each entry identifies its source and distinguishes a user need, an acceptance definition, a design decision or a technical constraint. Capacity limits describe the selected operating envelope; latency and sustained throughput require workload-specific measurements.')
    table(d,['ID','Question','Requirement / decision / constraint','Type','Basis'],[(a,b,c,types(t),s) for a,b,c,t,s in SOFTWARE],[1.2,3.2,6.7,2.2,2.6])
    h(d,'2. Functional Requirements (as DFDs & WNDs)',1,page=True)
    h(d,'2.1 Introduction',2)
    for text in FUNCTIONAL_INTRO:p(d,text)
    h(d,'2.2 Functional Requirements (FRs) - Data-Flow-Diagram Modelling',2,page=True)
    p(d,'External boxes represent document users, rounded nodes represent transformations, and named arrows show the data exchanged. Functional names identify the nodes; FR labels retain requirement traceability. The diagrams describe PDF data flows rather than deployment topology.')
    fig(d,'context','Figure 1. PDF Extractor context and input/output boundary.')
    table(d,['Boundary data','Meaning and constraint'],[
        ('Input document','One uploaded PDF copy; at most 50 MiB and 200 pages. The original is not overwritten.'),
        ('Operation parameters','Validated page selections, regular expressions, export choices or modification settings.'),
        ('Sensitive input','An opening password, signing container, container password or optional trust certificate used by the selected operation.'),
        ('Output','Structured results, a CSV file, an image ZIP package or a new PDF, depending on the operation.'),
        ('Feedback','A completion status, an explicit empty result or a stable error code with a readable explanation.')],[4,11.9])
    for i,g in enumerate(GROUPS,1):
        title,name,intro,branches,outputs,accept,extra=g
        h(d,f'2.2.{i} {title}',3,page=True);p(d,intro)
        fig(d,name,f'Figure {i+1}. {title}.')
        p(d,outputs);p(d,'Acceptance conditions: '+accept);p(d,extra)
    h(d,'2.3 Functional Requirements - Error Handling Common Use Package',2,page=True)
    p(d,'All tools share input, range, password, permission and output-storage validation. Feedback identifies the problem and the next usable action while preserving the source. A business conclusion, such as an invalid signature, is not the same as a failed HTTP request.')
    table(d,['Condition','Code / result','User-facing explanation'],ERRORS,[3.7,5.4,6.8],compact=True)
    p(d,'Additional shared limits return 413 for excessive size/pages/output and 429 CAPACITY_EXCEEDED when core processing capacity is occupied. An operation exceeding 90 seconds returns 504 PDF_TIMEOUT; a failed processing subprocess returns 503 without a partial result. Password and permission errors are separate. Client responses omit credentials, signing keys, stack traces and internal file paths.')
    h(d,'2.4 Functional Requirements - Window Navigation Diagram Prototype',2,page=True)
    p(d,'The interface uses a persistent document preview with operation panels. Selecting a tool keeps the current document; selecting a new file replaces it. One-based page selection, explicit save controls and local status/error feedback are shared conventions. Keyboard focus and readable controls apply across panels.')
    fig(d,'wnd','Figure 8. Window and panel navigation for the PDF workbench.')
    p(d,'The navigation covers file selection, document information, text, tables, images, encryption/security, signatures, output/save controls, and status/error feedback. Watermarking extends the same panel pattern. The following wireframe shows a text operation with its parameters and result beside the source preview.')
    fig(d,'text-prototype','Figure 9. English text-tool interface prototype.')
    h(d,'Swagger interface prototype',3)
    p(d,'The offline Swagger UI is available at /part-a/submission/swagger/. Its OpenAPI 3.0.3 contract is /part-a/submission/openapi.json, with the same-origin server /. It describes eleven core operations using multipart request bodies, typed result models and stable error responses. No credential has a default or example value. The WND represents user navigation; Swagger represents the associated request/response interface.')
    rows=[]
    for path,entry in spec['paths'].items():
        op=entry['post'];rows.append((op['summary'].split(' ')[0],path,{'W02':'Text and search','W03':'Tables and CSV','W04':'Images and ZIP','W05':'Document information','W06':'Encryption and permissions','W07':'Digital signatures','W08':'Watermarks'}[op['x-wnd']],', '.join(op['x-requirements'])))
    table(d,['ID','POST path','Panel','Requirement links'],rows,[1.2,6.4,3.7,4.6])
    shot=OUT/'diagrams/swagger-prototype.png'
    if shot.exists():fig(d,'swagger-prototype','Figure 10. English Swagger interface prototype served through the same-origin gateway.')
    p(d,'Reference basis: PDF Extractor project brief (FR1–FR19 and watermarking scope); Part A Student Template for Submission; the accompanying OpenAPI contract. API schemas and acceptance rules use the selected project scope recorded in sections 1.3.3 and 1.3.4.')
    path=OUT/'PartA_Submission.docx';d.save(path)
    return path

def build():
    OUT.mkdir(parents=True,exist_ok=True)
    for title,paragraphs in [('1.1',PROJECT_INTRO),('2.1',FUNCTIONAL_INTRO)]:
        n=len(re.findall(r"\b[\w]+(?:[-’'][\w]+)*\b",' '.join(paragraphs)))
        assert 200<=n<=300,(title,n)
    diagrams();spec=contract();path=document(spec)
    (OUT/'index.html').write_text('''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>PDF Extractor Part A</title><style>body{max-width:900px;margin:48px auto;padding:0 24px;font:17px/1.8 Arial;color:#222}a{color:#222}li{margin:12px 0}</style><h1>PDF Extractor Part A</h1><p>Requirements Specification and Design</p><ul><li><a href="PartA_Submission.docx">English submission document</a></li><li><a href="swagger/">Interactive Swagger interface prototype</a></li><li><a href="openapi.json">OpenAPI 3.0.3 contract</a></li><li><a href="diagrams/wnd.svg">Window navigation diagram</a></li></ul></html>''',encoding='utf-8')
    print(json.dumps({'document':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'figures':9+int((OUT/'diagrams/swagger-prototype.png').exists())}))
if __name__=='__main__':build()
